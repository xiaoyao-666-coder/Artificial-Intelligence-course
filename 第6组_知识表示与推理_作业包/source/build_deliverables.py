"""使用 Codex bundled Python 生成报告、可编辑 PPT 和视频讲稿。"""
from pathlib import Path
from copy import deepcopy
import hashlib, json, zipfile
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pptx import Presentation
from pptx.util import Inches as PI, Pt as PP
from pptx.dml.color import RGBColor as PC
from PIL import Image, ImageDraw, ImageFont

ROOT=Path(__file__).parent
OUT=ROOT/'deliverables'; OUT.mkdir(exist_ok=True)
QA=ROOT/'qa'; QA.mkdir(exist_ok=True)
bench=json.loads((ROOT/'output/benchmark.json').read_text(encoding='utf-8'))
timing={r['method']:r['median_ms'] for r in bench['results'] if r['depth']==8}
speed=timing['SQL 递归 CTE']/timing['图 DFS']
title='企业股权知识表示与风险推理'
sections={
'摘要': '本项目面向企业股权穿透和关联风险核查，研究知识表示形式如何影响推理过程与查询效率。系统以自制的4名自然人、12家企业和25条关系为输入，联合使用知识图谱、谓词事实与产生式规则、框架表示三种方法。图谱承担多跳路径与环路查询，15条规则完成控制关系和风险线索推导，企业框架承接继承信息及推理结果。实现流程为数据建模、模式校验、事实转换、正向推理、证据解释、图计算和知识回写。程序由53条事实扩充到135条，识别循环持股、担保圈及关联交易待核查线索，11项功能测试全部通过。另以430个实体的合成网络比较图遍历与两类SQL查询，检查结果一致后计时。结果说明，表示方式应服务于任务；控制权不等于穿透持股，规则结论也不等于法律认定。',
'背景与目标': '企业尽职调查往往需要回答谁控制一家企业、交易双方是否存在共同控制、担保责任能否沿链条传播等问题。单独查看工商登记表，只能看到局部股东关系；多层持股、关联交易和担保混合后，人工核查容易遗漏跨企业联系。同名实体、历史变更和缺失资料还会进一步增加核查难度。\n真实行业案例选用ICIJ Offshore Leaks公开数据库。其说明页展示企业、人员及中介之间的联系，指出部分企业又是其他企业的股东，并提醒数据可能重复、发生变化，入库不代表违法[1]。这一公开实践说明，将分散记录组织为实体关系网络，有助于追溯公司背后的人。\n本项目将其关系追踪思路缩小为课堂案例，不声称复现ICIJ内部系统，也未使用其真实个人记录。目标是对比至少两类知识表示，演示可解释推理及知识增强，并用等价查询考察效率。系统输出风险线索与证据路径，最终判断仍需核查原始材料。',
'相关技术与工具准备': '知识图谱以实体、关系和属性组织知识。本项目以NetworkX有向多重图存储人员和企业，允许同一对实体同时存在持股、担保或交易关系；模式层规定各类边的起止类型和属性。它适合多跳追踪和环路检测，但实体对齐、时间版本及证据质量需要额外治理。\n谓词把关系写成Holds(P1,C1,0.8)等事实，产生式规则采用IF前提THEN结论。推理机通过变量绑定和合一匹配规则，按优先级消解冲突，正向扩展事实库，记录每条新结论的来源。该表示适合明确、可审计的业务规则；规则数量和匹配组合增大时成本会上升，缺少事实不能直接当作否定。\n框架通过企业、上市公司和担保公司之间的is_a继承共享槽值，适合企业档案及缺省属性管理。它比图谱更方便组织单个对象的属性，但不擅长全局路径计算。缺省风险等级只是程序初始状态，不是对企业安全性的评估。\n实现使用Python 3.13.1、NetworkX 3.5、SQLite 3.45.3、Matplotlib及可选PyVis。本项目没有模型训练和GPU需求。课堂数据全部虚构，CSV以UTF-8 BOM导出；效率实验使用固定随机种子42生成分层网络，便于重复核对。',
'算法设计与实现': '系统流程为：构建数据→模式校验→图谱转谓词→正向规则推理→路径和环路分析→结论回写框架及图谱。实体通过稳定ID关联，持股比例使用0到1的小数，担保与交易金额单位为万元。图谱原始边共25条；实体类型事实、企业分类事实与边事实合计53条。\n规则库包括控制识别R1—R6、共同控制交易与可达关系R7—R11、风险线索R12—R14和教学风险定级R15。每轮构造可产生新事实的冲突集，优先执行高优先级规则；同优先级按规则顺序选择，执行该规则的所有匹配实例。集合去重防止重复加入事实，无新增事实时终止。解释模块沿结论记录递归回溯已知事实。\n例如，王建国持有华鑫控股80%，后者持有华鑫实业65%；华鑫实业与华鑫控股分别持有恒通科技35%和20%。在本项目的简化控制规则下，两条可支配股权合计55%，于是推出华鑫控股控制恒通科技，再传递得到王建国的控制关系。王建国和配偶持有远航贸易30%与25%；教学上假设二者一致行动，得到共同控制交易线索。配偶身份本身不能代替真实控制协议。\n经济持股按路径比例连乘、不同简单路径相加：王建国对恒通科技的比例为0.8×0.2＋0.8×0.65×0.35＝34.2%。该结果与上述控制关系含义不同。DFS向上遍历股东，维护已访问节点，遇环截断；深度截断处的企业不会被标记为自然人候选。25%是本程序的教学筛查阈值，不是完整法律标准。\n环路分析分别抽取持股和担保子图，通过有向简单环识别循环；关联路径查询可忽略方向寻找最短联系，但返回步骤保留原始箭头。最后把9条自然人实际控制关系写回图谱，并给6家企业写入风险标签，框架更新风险提示。原始事实与推断结论通过关系类型及来源规则区分，便于解释和复核。',
'测试结果与分析': f'测试在Windows 11、Python 3.13.1的CPU环境运行，无训练阶段。11项自动化测试覆盖控制链、34.2%计算、三类预警、图与规则环路一致性、负例、框架继承、深度截断、50%边界、重复变量合一、推理不动点和SQL等价性。全部通过仅代表这些设计用例，不代表真实风控准确率。演示产生17个推理周期，事实数由53增至135；恒通科技出现关联交易待核查线索，循环持股与担保圈各涉及3家企业，因鼎丰建设重叠，共6家企业带标签。\n效率实验使用430个实体、792条持股边的无环网络，查询深度为2、4、6、8。每个深度固定20个目标，先逐一校验三种方法的股东、比例及深度结果多重集合一致，随后重复9轮，报告每次查询平均耗时的轮间中位数。SQLite对被投企业建索引；计时包含结果物化，排除建库和绘图，建库耗时另存JSON。\n8层查询中，图DFS为{timing["图 DFS"]:.3f}毫秒，递归CTE为{timing["SQL 递归 CTE"]:.3f}毫秒，逐层JOIN为{timing["SQL 逐层 JOIN"]:.3f}毫秒。该实验支持图遍历在当前内存查询中的优势；硬件、缓存、查询顺序和数据分支数会影响结果，不能推广为生产图数据库的普遍性能结论。',
'问题与改进思路': '接续检查发现，原代码把共同控制交易称为隐性关联交易，但数据并不包含披露状态，无法据此判断隐瞒行为，因此已改为关联交易待核查。同时，原先任意亲属均触发一致行动，现在仅保留配偶这一明确的教学假设，并在规则中注明实务需要另行取证。深度截断时还需区分企业中间节点与最终自然人，避免把企业错误标记为筛查候选。\n现有控制规则只覆盖单个子公司与自身合并持股，不完整处理多主体共同控制、同股不同权及复杂协议。循环股权采用简单路径截断，会遗漏无限往返产生的经济权益贡献，不能代替矩阵求解。框架缺省披露与风险字段也只是课堂示例，不承诺适用于所有法规和例外。\n后续可加入日期、证据来源、身份消歧和明确的一致行动协议；规则变化时通过依赖图撤回失效结论；对大规模图使用索引、增量推理和查询规划。概率推理可表达不确定性，大模型可辅助提取候选关系，但仍需结构校验和人工复核。',
'总结': '本项目完成了企业股权场景下从知识表示到推理解释、再到知识增强的闭环。图谱表达实体关系，产生式系统推导可追溯结论，框架整理企业属性；三者在同一数据上协作，使局部股东记录转化为可检查的控制链与风险线索。效率实验也表明，相同知识采用不同组织和查询方法，会产生不同执行成本。\n实践中最关键的区分是经济持股、规则控制和法律判断：34.2%的穿透持股可以与简化规则中的控制关系同时成立，但二者均不能自动构成违法认定。当前成果是可重复运行的小型教学系统，尚不具备真实业务部署所需的数据完整性、规则覆盖和动态更新能力。未来应以证据质量、语义一致性和可解释性为前提扩展规模。'
}
refs=[
'[1] ICIJ. About the Offshore Leaks Database[EB/OL]. https://offshoreleaks.icij.org/pages/about, 访问日期2026-09-28.',
'[2] HOGAN A, BLOMQVIST E, COCHEZ M, et al. Knowledge Graphs[J]. ACM Computing Surveys, 2021, 54(4): Article 71. DOI:10.1145/3447772.',
'[3] RUSSELL S, NORVIG P. Artificial Intelligence: A Modern Approach[M]. 4th ed. Pearson, 2021.',
'[4] MINSKY M. A Framework for Representing Knowledge[R]. MIT AI Laboratory, Memo 306, 1974.',
'[5] NetworkX Developers. NetworkX Documentation[EB/OL]. https://networkx.org/documentation/stable/, 访问日期2026-09-28.',
'[6] SQLite. The WITH Clause[EB/OL]. https://www.sqlite.org/lang_with.html, 访问日期2026-09-28.',
'[7] 课程教学资料. 第2章 my知识表示与知识图谱[PPT]. 项目目录所附课件.'
]
reference=next(ROOT.parent.glob('*.docx'))
doc=Document(reference)
pars=doc.paragraphs
# Retain the template paragraph patterns and original package; only body slots change.
body_pattern=deepcopy(pars[17]._p); heading_pattern=deepcopy(pars[16]._p)
contract={'reference':str(reference),'sha256':hashlib.sha256(reference.read_bytes()).hexdigest(),'reference_pages':2,'sections':1,'page_cm':[21.00086,29.70036],'margins_cm':[doc.sections[0].top_margin.cm,doc.sections[0].bottom_margin.cm,doc.sections[0].left_margin.cm,doc.sections[0].right_margin.cm],'slots':'word/document.xml body paragraphs 0..31: replace instructions with matching ten report sections; clone heading and body paragraph patterns for expanded content. Preserve sectPr and all other template parts except necessary image relationships. No tables or controls in source.','fonts':'Source Normal 10.5pt, source heading bold; title 14pt; preserve source paragraph properties, black text.','render':'qa/template.pdf and template-1.png, template-2.png; Word fallback because packaged renderer lacks LibreOffice.'}
(QA/'artifact.md').write_text('# Template contract\n'+json.dumps(contract,ensure_ascii=False,indent=2),encoding='utf-8')
for p in list(doc.paragraphs): p._p.getparent().remove(p._p)
def paragraph(text,heading=False):
    from docx.text.paragraph import Paragraph
    e=deepcopy(heading_pattern if heading else body_pattern)
    for child in list(e):
        if child.tag!=qn('w:pPr'): e.remove(child)
    doc._body._body.insert(len(doc._body._body)-1,e)
    p=Paragraph(e,doc._body);r=p.add_run(text)
    r.font.name='Times New Roman';r.font.size=Pt(10.5);r.font.bold=heading;r.font.color.rgb=RGBColor(0,0,0)
    r._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'宋体')
    p.paragraph_format.keep_with_next=heading;p.paragraph_format.widow_control=True
    return p
p=paragraph('人工智能课堂报告',True);p.runs[0].font.size=Pt(14);p.alignment=1
paragraph('第6组    姓名与学号：待填写    成员分工：待填写')
paragraph('一、标题',True);paragraph(title)
heading_names=['二、摘要','三、关键词','四、背景与目标','五、相关技术与工具准备','六、算法设计与实现','七、测试结果与分析','八、问题与改进思路','九、总结','十、参考文献']
for heading in heading_names:
    paragraph(heading,True)
    key=heading.split('、',1)[1]
    content='知识图谱；产生式推理；框架表示；股权穿透；风险核查' if key=='关键词' else ('\n'.join(refs) if key=='参考文献' else sections[key])
    for block in content.split('\n'): paragraph(block)
    if key=='算法设计与实现':
        paragraph('核心代码示例  股权路径比例计算',True)
        paragraph('r = ratio * H[p][node]["ratio"]\nrows.append((p, r, depth + 1))\ndfs(p, r, depth + 1, path | {p})')
        paragraph('摘自 reasoning/graph_reasoning.py。每走一层累乘比例，并记录路径中的节点以避免重复访问。')
    if key=='测试结果与分析':
        for depth in (2,4,6,8):
            values=[r for r in bench['results'] if r['depth']==depth]
            paragraph(f'{depth}层查询  '+ '；'.join(f'{r["method"]} {r["median_ms"]:.3f} ms' for r in values))
        p=paragraph('');p.add_run().add_picture(str(ROOT/'output/kg_reasoned.png'),width=Inches(5.65))
        paragraph('图1 推理增强后的知识图谱  灰线为原始关系，实控边标注来源规则，红圈表示教学风险线索。')
doc.save(OUT/'课堂报告.docx')
# Keep every untouched OOXML part byte-identical to the source template.
target=OUT/'课堂报告.docx'
with zipfile.ZipFile(reference) as src, zipfile.ZipFile(target) as dst:
    parts={n:dst.read(n) for n in dst.namelist()}
    for n in src.namelist():
        if n not in ('word/document.xml','word/_rels/document.xml.rels','[Content_Types].xml'): parts[n]=src.read(n)
with zipfile.ZipFile(target,'w',zipfile.ZIP_DEFLATED) as z:
    for n,data in parts.items(): z.writestr(n,data)

slides=[
('企业股权知识表示\n与风险推理','课堂报告 1  /  第6组',[('16','实体'),('25','原始关系'),('15','推理规则')], '我们选择企业股权风控场景，研究知识如何表示、如何推理，以及不同表示对查询效率的影响。系统用十六个虚构实体完成一个可以运行和解释的小型案例。'),
('从分散登记到关系追踪','真实案例  ICIJ Offshore Leaks',['公开案例：人员、企业和中介构成关系网络','行业流程：汇集记录 → 实体对齐 → 关系查询 → 人工核查','课堂实现：使用虚构数据，演示股权与担保推理'], '真实案例参考国际调查记者同盟的离岸数据库。它把企业、人员和中介关联起来，帮助追踪公司背后的人。我们借鉴关系追踪思路，用虚构数据演示，入库和存在关联都不代表违法。'),
('三种表示各有分工','表示方式要服务于具体任务',['知识图谱｜实体与边：适合多跳、路径和环路','谓词与产生式｜IF–THEN：适合明确规则及证据解释','框架｜槽与继承：适合企业档案及缺省属性'], '知识图谱适合多跳路径和环路查询，产生式规则适合明确的业务条件和结论解释，框架通过继承组织企业属性。图谱需要治理实体质量，规则面临匹配成本，框架则不擅长全局关系计算。'),
('控制权 ≠ 穿透持股','同一组事实，回答两个不同问题',['教学控制链：80% → 65% → 合并35%＋20%','经济持股：80%×20%＋80%×65%×35%＝34.2%','配偶一致行动为教学假设，实务需要额外证据'], '王建国持有华鑫控股百分之八十。华鑫控股控制华鑫实业，后两者合计持有恒通科技百分之五十五，因此教学规则推出控制关系。但经济持股沿路径相乘再求和，只有百分之三十四点二。两者不能混淆。'),
('推理结论可以追溯','匹配 → 冲突消解 → 执行 → 不动点',['共同控制人 ＋ 企业间交易 → 关联交易线索','上市公司参与 → 关联交易待核查预警','53条初始事实 → 17个周期 → 135条事实'], '正向推理机先匹配事实，再按优先级执行规则，直到不能产生新结论。当前运行经过十七个周期，事实从五十三条增加到一百三十五条。每个结论都能回溯到来源规则和原始事实，交易预警只表示需要核查。'),
('知识增强形成闭环','推理回写，原始事实与推断结论分开','graph', '系统把九条自然人控制关系写回图谱，给六家企业标记风险线索，同时更新企业框架。循环持股和担保圈各涉及三家企业，另有一家上市公司的交易预警，因为存在重叠，共有六家企业带标签。'),
('效率比较必须先保证等价','430实体 · 792条边 · 20个目标 · 9轮重复',[f'8层 图DFS：{timing["图 DFS"]:.3f} ms',f'8层 SQL递归CTE：{timing["SQL 递归 CTE"]:.3f} ms',f'8层 SQL逐层JOIN：{timing["SQL 逐层 JOIN"]:.3f} ms','相同路径结果；建库与绘图不计入查询耗时'], '效率实验使用固定的无环网络，先检查三种方法的查询结果一致，再重复九轮计时。八层查询中，图遍历约零点五六毫秒，递归查询约一点零三毫秒，逐层连接约零点七七毫秒。结果仅适用于当前数据和实现。'),
('可靠演示与实际部署之间','11项测试通过，仍有明确边界',['功能验证：控制链、环路、负例、边界及查询等价','当前局限：合成数据、简单控制规则、环路截断','下一步：时间与来源、身份消歧、增量推理、人工复核'], '十一项测试全部通过，但这不等于真实风控准确率。当前模型仍使用合成数据、简化控制规则和环路截断。后续需要增加时间与证据来源、身份消歧，以及规则更新后的结论撤回机制。'),
('用合适的表示支持可解释推理','结论与参考资料',['图谱找关系，规则推结论，框架组织属性','真实案例：ICIJ Offshore Leaks About 页面','理论：Hogan等 Knowledge Graphs；Russell与Norvig','工具：NetworkX、SQLite；完整文献见报告'], '本项目展示了知识表示、推理解释和知识增强的完整过程。核心结论是，表示方式要与任务匹配；推理结果要保留证据和适用条件。报告附有完整参考文献，代码、演示视频和效率数据可以复查。')
]
prs=Presentation();prs.slide_width=PI(13.333);prs.slide_height=PI(7.5)
navy='142E43';teal='087F8C';ink='203746';muted='587080'
def textbox(s,x,y,w,h,text,size=24,color=ink,bold=False):
    box=s.shapes.add_textbox(PI(x),PI(y),PI(w),PI(h));tf=box.text_frame;tf.word_wrap=True
    for i,line in enumerate(text.split('\n')):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph();p.text=line;p.font.name='Microsoft YaHei';p.font.size=PP(size);p.font.bold=bold;p.font.color.rgb=PC.from_string(color);p.space_after=PP(14)
    return box
for i,(head,sub,items,narration) in enumerate(slides):
    s=prs.slides.add_slide(prs.slide_layouts[6]);s.background.fill.solid();s.background.fill.fore_color.rgb=PC.from_string('F6F8FA')
    textbox(s,.65,.35,12,.4,f'知识表示与推理  /  EQUITY KNOWLEDGE GRAPH',12,teal,True)
    textbox(s,.65,1.0,12,1.4 if i==0 else .8,head,34,navy,True)
    textbox(s,.7,2.5 if i==0 else 1.95,12,.6,sub,18,muted)
    if i==0:
        for j,(num,label) in enumerate(items):
            textbox(s,.8+j*4,4,3,1,num,50,teal,True);textbox(s,.8+j*4,5.05,3,.5,label,21,muted)
    elif items=='graph':
        s.shapes.add_picture(str(ROOT/'output/kg_reasoned.png'),PI(3.0),PI(2.7),height=PI(4.2))
    else:
        for j,item in enumerate(items): textbox(s,.85,2.85+j*.85,11.7,.77,item,23,ink)
    textbox(s,.7,7.05,10,.25,'课堂演示数据为虚构；推理输出是核查线索。',10,muted)
    textbox(s,12.05,7.0,.6,.3,f'{i+1:02}',12,teal)
    s.notes_slide.notes_text_frame.text=narration
prs.save(OUT/'课堂汇报.pptx')
(QA/'narration.json').write_text(json.dumps([x[3] for x in slides],ensure_ascii=False,indent=2),encoding='utf-8')
(OUT/'汇报讲稿.md').write_text('# 课堂汇报讲稿\n\n'+'\n\n'.join(f'## 第{i+1}页 {s[0].replace(chr(10), " ")}\n{s[3]}' for i,s in enumerate(slides)),encoding='utf-8')
(QA/'report_content.json').write_text(json.dumps(sections,ensure_ascii=False,indent=2),encoding='utf-8')
print('Created report, editable PPT, and narration script.')
import runpy
runpy.run_path(str(ROOT/'revise_ppt.py'))
