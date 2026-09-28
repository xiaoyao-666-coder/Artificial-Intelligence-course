"""Revise an existing report without overwriting its source template or original deliverable."""
from pathlib import Path
from docx import Document
from docx.shared import Pt
ROOT=Path(__file__).resolve().parent
source=ROOT/'deliverables/课堂报告.docx'
out=ROOT/'deliverables/课堂报告_修订版.docx'
doc=Document(source)
replacements={
 '例如，王建国': '例：王建国持华鑫控股80%，华鑫控股持华鑫实业65%；二者分别持恒通科技20%和35%。简化规则合并可支配股份为55%，推出华鑫控股及王建国的控制关系。王建国与配偶另持远航贸易30%和25%，依教学一致行动假设推出关联交易线索；配偶关系不能代替实际控制协议。',
 '规则库包括控制识别': '规则库共15条，覆盖控制识别、共同控制交易、可达关系、环路及教学风险定级。每轮合一匹配规则前提，形成冲突集；优先执行高优先级规则，同优先级按规则顺序选择，执行其全部匹配实例。事实集合去重，无新增事实即终止；解释模块沿推导记录回溯证据。',
 '测试在Windows 11': '原基准实验运行于Windows、Python 3.13.1的CPU环境，无模型训练和GPU阶段。当前环境已独立复跑11项功能测试，均通过，覆盖控制链、34.2%经济持股、预警、框架继承、深度截断、50%边界、合一与不动点、SQL等价性。主程序复跑得到53→135条事实、17个周期、9条自然人控制关系及6家风险标记企业；这些是教学用例结果，不代表真实风控准确率。',
 '效率实验使用430个实体': '原效率实验固定种子42，使用430实体、792条持股边的DAG。对20个目标、2/4/6/8层逐一检查DFS、递归CTE和逐层JOIN的结果多重集合相同；随后重复9轮，取每次查询平均耗时的轮间中位数。计时含结果物化，不含建库和绘图；SQLite已建索引。',
 '8层查询中': '下列为仓库保留的原基准结果，不冒充本次机器的新测量。它比较内存图算法与SQLite实现，不代表全部图数据库优劣；硬件、缓存、分支数及查询顺序均可能影响结果。',
}
for p in doc.paragraphs:
 for prefix,text in replacements.items():
  if p.text.startswith(prefix):
   p.text=text
   for r in p.runs:
    r.font.name='宋体';r.font.size=Pt(10.5)
   break
# Preserve all ten sections, member placeholders, references and original figure.
# Source generator cloned paragraph IDs; Word expects unique IDs for layout/edit tracking.
for i, paragraph in enumerate(doc.paragraphs, 1):
    paragraph._p.set('{http://schemas.microsoft.com/office/word/2010/wordml}paraId', f'{i:08X}')
    paragraph._p.set('{http://schemas.microsoft.com/office/word/2010/wordml}textId', f'{i+1000:08X}')
doc.save(out)
print(out)
