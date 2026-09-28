"""按背景—问题—案例—方法—验证—结论重写课堂PPT。"""
from pathlib import Path
import json
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

ROOT=Path(__file__).parent; OUT=ROOT/'deliverables'; QA=ROOT/'qa'
prs=Presentation();prs.slide_width=Inches(13.333);prs.slide_height=Inches(7.5)
NAVY='18364A';INK='294455';TEAL='087F8C';GRAY='607686';BG='F5F7FA'
notes=[];titles=[]
def text(s,x,y,w,h,value,size=22,color=INK,bold=False):
    sh=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h));tf=sh.text_frame;tf.word_wrap=True
    for i,line in enumerate(value.split('\n')):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph();p.text=line;p.font.name='Microsoft YaHei';p.font.size=Pt(size);p.font.bold=bold;p.font.color.rgb=RGBColor.from_string(color);p.space_after=Pt(9)
    return sh
def rect(s,x,y,w,h,color='FFFFFF'):
    sh=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h));sh.fill.solid();sh.fill.fore_color.rgb=RGBColor.from_string(color);sh.line.fill.background();return sh
def card(s,x,y,w,h,label,body):
    rect(s,x,y,w,h);text(s,x+.2,y+.15,w-.4,.55,label,22,TEAL,True);text(s,x+.2,y+.85,w-.4,h-.95,body,19)
def slide(section,title,sub,note):
    s=prs.slides.add_slide(prs.slide_layouts[6]);s.background.fill.solid();s.background.fill.fore_color.rgb=RGBColor.from_string(BG)
    text(s,.55,.22,12,.4,section,13,TEAL,True);text(s,.55,.83,12.2,.8,title,32,NAVY,True);text(s,.6,1.75,12,.7,sub,19,GRAY)
    text(s,.6,7.08,11,.23,'人工智能课堂报告 · 第6组 · 企业股权知识表示与风险推理',10,GRAY)
    text(s,12,7.02,.7,.35,f'{len(prs.slides):02}',13,TEAL)
    s.notes_slide.notes_text_frame.text=note;notes.append(note);titles.append(title);return s

s=slide('选题介绍','企业股权知识表示与风险推理','从“谁持有股份”出发，追踪“谁控制企业、哪些关系需要核查”。','本次汇报研究企业股权知识表示与风险推理。我们从一个实际问题出发：银行或投资者审查一家企业时，如何从分散的股东记录中找到背后的控制人，并发现需要核查的关联关系？')
text(s,.8,2.9,11.7,1.15,'当企业背后还有企业，\n只看一张股东表，就很难看清完整关系。',30,NAVY,True)
rect(s,.8,4.65,11.7,1.15,'E1EFF1');text(s,1.05,4.9,11.2,.75,'汇报主线：业务背景 → 案例问题 → 知识表示 → 推理与验证',23,TEAL,True)
text(s,.85,6.25,11.5,.45,'组员姓名、学号与分工：待填写',16,GRAY)

s=slide('01 研究背景','为什么需要穿透企业股权关系？','股权穿透：沿着多层股东关系向上追踪，找到最终的自然人及其持股路径。','企业股权穿透，是沿着多层股东关系向上追踪。在贷款审批和投资尽调中，审查人员既要知道谁控制企业，也要判断交易双方是否有关联。多层持股、交叉担保和分散资料，使人工核查容易遗漏关系。')
card(s,.65,2.8,3.9,2.65,'业务场景','银行贷款审批、投资尽调\n审查目标企业及其关联方')
card(s,4.72,2.8,3.9,2.65,'人工核查的困难','股东也可能是另一家企业\n交易与担保分散在不同记录中')
card(s,8.8,2.8,3.9,2.65,'希望得到的答案','谁在背后控制企业？\n风险线索由哪些事实支持？')
text(s,.85,5.95,11.8,.65,'因此，需要把分散记录连接起来，再按照明确规则追踪和推理。',24,NAVY,True)

s=slide('02 行业案例与本项目目标','从真实关系追踪案例缩小到课堂任务','行业参考：ICIJ Offshore Leaks 公开数据库连接人员、企业与中介信息。','真实行业案例参考国际调查记者同盟的离岸数据库。它连接人员、企业和中介，帮助追踪公司背后的人。本项目借鉴这种关系追踪思路，使用虚构数据，完成控制人识别、关联交易核查和环路检测三个课堂任务。')
card(s,.7,2.8,5.75,2.85,'真实案例给出的启发','企业可以是其他企业的股东。\n把记录组织成关系网络，才能沿链追踪。\n数据仍需身份核实，出现关联不等于违法。')
card(s,6.7,2.8,5.9,2.85,'我们的可运行任务','① 识别多层股权背后的控制人\n② 找出共同控制下的交易线索\n③ 检测循环持股与担保圈')
text(s,.85,5.98,11.7,.6,'课堂范围：4名自然人、12家虚构企业、25条原始关系；不复现ICIJ内部系统。',19)
text(s,.85,6.58,11.5,.28,'来源：https://offshoreleaks.icij.org/pages/about',12,GRAY)

s=slide('03 案例与问题','以恒通科技为例：多层股权怎样影响判断？','以下名称与比例均来自本项目的虚构数据；箭头方向表示“股东 → 被投资企业”。','来看贯穿演示的恒通科技案例。王建国持有华鑫控股百分之八十，华鑫控股持有华鑫实业百分之六十五。后二者分别持有恒通科技百分之二十和三十五。现在的问题是：王建国是否控制恒通科技，经济持股又是多少？')
def node(x,y,w,label):
    rect(s,x,y,w,.72,'E1EFF1');text(s,x+.07,y+.12,w-.14,.47,label,22,TEAL,True)
node(.95,2.9,2.2,'王建国');node(4.1,2.9,2.4,'华鑫控股');node(8.5,2.9,2.5,'华鑫实业');node(7.05,5.05,2.8,'恒通科技')
text(s,3.18,3.05,.95,.5,'80%→',16);text(s,6.65,3.05,1.7,.5,'65% →',20)
text(s,5.05,4.02,2.9,.55,'20%  ↘',24);text(s,9.18,4.02,2.3,.55,'↙  35%',24)
text(s,.9,5.05,5.5,1.1,'问题一：谁能够控制恒通科技？\n问题二：沿股权路径计算的权益是多少？',21,NAVY,True)
text(s,.9,6.35,11.7,.45,'还需继续核查：若它与同一控制人旗下的企业交易，是否存在关联交易线索？',20)

s=slide('04 知识表示方法','把同一批记录表示成三种可计算的知识','图谱负责关系查询，规则负责条件推导，框架负责企业档案；三者共同支撑案例分析。','为回答这些问题，我们采用三种表示。图谱把人和企业表示为节点，适合找路径。谓词与产生式把持股事实和控制条件写成规则，便于解释结论。框架通过槽和继承整理企业档案，但不擅长全局路径计算。')
rows=[('知识图谱','王建国 —持股80%→ 华鑫控股','适合路径与环路查询','难点：实体对齐与数据质量'),('谓词与产生式','Holds(P1,C1,0.8) → 控制规则','适合明确条件与证据解释','难点：规则多时匹配成本增加'),('框架表示','恒通科技 is_a 上市公司','适合属性继承与档案更新','难点：缺省值不等于已知事实')]
for j,(a,b,c,d) in enumerate(rows):
    y=2.7+j*1.27;rect(s,.7,y,11.9,1.1);text(s,.9,y+.14,2.25,.65,a,23,TEAL,True);text(s,3.1,y+.1,5.05,.8,b+'\n'+c,18);text(s,8.35,y+.17,4,.7,d,18,GRAY)
text(s,.9,6.65,11.7,.3,'转换流程：关系记录 → 图谱与谓词事实 → 规则推导 → 结果写回企业框架和图谱',16)

s=slide('05 推理演示','沿着案例推出结论，并保留证据链','先说明控制规则，再计算经济持股；两者回答的问题不同。','在简化规则下，华鑫控股控制华鑫实业，能够合并支配恒通科技百分之五十五的股份，再向上推导王建国的控制关系。经济持股则按路径相乘后求和，得到百分之三十四点二。控制权和经济持股不能混为一谈。')
card(s,.7,2.7,5.8,2.45,'控制关系的推导','华鑫控股控制华鑫实业（65%）\n可支配股权：20%＋35%＝55%\n控制关系向上传递到王建国')
card(s,6.75,2.7,5.85,2.45,'经济持股的计算','路径一：80%×20%＝16%\n路径二：80%×65%×35%＝18.2%\n两条路径合计：34.2%')
rect(s,.7,5.45,11.9,1.13,'E1EFF1');text(s,.9,5.62,11.45,.85,'继续推理：共同控制人＋企业间交易 → 关联交易待核查\n每条结论记录来源规则和前提；配偶一致行动等属于教学假设，需另行取证。',19)

s=slide('06 系统运行结果','从原始关系得到可解释的风险线索','完整流程：输入记录 → 正向推理与图查询 → 解释依据 → 回写图谱及企业档案。','程序运行后，五十三条初始事实扩充为一百三十五条，九条自然人控制关系写回图谱。循环持股、担保圈和交易待核查共涉及六家企业。十一项功能测试通过。这些结果验证了演示逻辑，不代表真实业务准确率。')
s.shapes.add_picture(str(ROOT/'output/kg_reasoned.png'),Inches(.7),Inches(2.65),height=Inches(3.88))
card(s,6.1,2.7,6.5,3.95,'运行结果与验证','53条事实 → 17个周期 → 135条事实\n新增9条自然人控制关系\n6家企业带风险线索，支持追溯依据\n11项测试通过：控制链、环路、负例、边界\n图中黑边为推理关系，红圈为风险线索')

s=slide('07 效率实验','相同查询任务，不同表示的效率如何？','在相同无环数据上先核对查询结果一致，再比较图DFS、SQL递归CTE和逐层JOIN。','我们还比较了同一多跳股东查询的效率。固定四百三十个实体，查询二十个目标，重复九轮。八层查询中，图遍历约零点五六毫秒，SQL两种实现分别约一点零三和零点七七毫秒。优势仅适用于当前数据与实现。')
b=json.loads((ROOT/'output/benchmark.json').read_text(encoding='utf-8'));vals=[r for r in b['results'] if r['depth']==8]
for j,r in enumerate(vals):
    y=2.85+j*.88;text(s,.9,y,3.1,.55,r['method'],23,NAVY,True);rect(s,4.1,y+.03,r['median_ms']*4.7,.4,TEAL if j==0 else '7894A8');text(s,9.5,y-.03,2.7,.65,f'{r["median_ms"]:.3f} ms',24,TEAL,True)
text(s,.95,5.65,11.5,.65,'430实体 / 792条边 / 20个目标 / 9轮重复；显示8层查询的轮间中位耗时。',18)
text(s,.95,6.25,11.5,.5,'计时包含结果物化，不含建库与绘图；小规模内存实验不等同于生产数据库评测。',18,GRAY)

s=slide('08 总结与展望','回到业务问题：把分散记录变成有依据的线索','本项目完成了一个小型、可运行、可追溯的知识表示与推理系统。','回到开头的问题，系统能够把分散股东记录连接起来，推导控制链，并给出关联风险的证据。图谱、规则和框架各有所长。当前仍受虚构数据和简化规则限制，下一步需要补充时间与来源、身份核验和动态更新，再用于真实核查。')
card(s,.7,2.75,5.8,2.7,'本次完成了什么','用三种知识表示组织同一案例\n演示控制关系、路径和环路推理\n用证据链与等价查询验证结果')
card(s,6.75,2.75,5.85,2.7,'仍有哪些不足','虚构数据，控制规则覆盖有限\n循环股权按简单路径截断\n后续加入时间、来源、身份核验及增量推理')
text(s,.9,5.85,11.6,.55,'结论：表示方法要匹配任务；推理结论要有证据，也要说明适用条件。',23,NAVY,True)
text(s,.9,6.55,11.6,.4,'参考：ICIJ Offshore Leaks；Hogan等《Knowledge Graphs》；课程第2章；NetworkX与SQLite文档。完整书目见报告。',11,GRAY)

prs.save(OUT/'课堂汇报.pptx')
(QA/'narration.json').write_text(json.dumps(notes,ensure_ascii=False,indent=2),encoding='utf-8')
(OUT/'汇报讲稿.md').write_text('# 课堂汇报讲稿\n\n'+'\n\n'.join(f'## 第{i+1}页 {title}\n{note}' for i,(title,note) in enumerate(zip(titles,notes))),encoding='utf-8')
print('Revised 9 slides with background, case, methods, results and conclusion.')
