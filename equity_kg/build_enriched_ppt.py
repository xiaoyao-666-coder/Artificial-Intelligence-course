"""Build the editable 14-slide warm, saturated project presentation.

Uses the published revised deck only to preserve its complete team notes.
Performance figures are read from the historical benchmark; no new timing claims.
Run with the Codex bundled Python (python-pptx and Pillow).
"""
from pathlib import Path
import json
import math

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR, MSO_AUTO_SIZE
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_TICK_MARK
from pptx.chart.data import CategoryChartData
from pptx.oxml.xmlchemy import OxmlElement

ROOT = Path(__file__).resolve().parent
OUT = ROOT / 'deliverables'
QA = ROOT / 'qa' / 'ppt-enriched'
QA.mkdir(parents=True, exist_ok=True)
OUT.mkdir(exist_ok=True)
SOURCE = OUT / '课堂汇报_修订版.pptx'
TARGET = OUT / '企业股权知识表示与风险推理_暖色增强版.pptx'
old = Presentation(SOURCE)
team_notes = old.slides[0].notes_slide.notes_text_frame.text.split('组员与分工', 1)[1].strip()
benchmark = json.loads((ROOT / 'output/benchmark.json').read_text(encoding='utf-8'))

P = Presentation()
P.slide_width, P.slide_height = Inches(13.333333), Inches(7.5)
C = dict(bg='F5EBDD', paper='FFF9EF', red='8F2843', wine='551C30', ink='252526',
         orange='D9652F', gold='F2BD58', muted='79665D', light='E8D1B4', white='FFFFFF',
         green='326759', line='D7BFA6', pink='B85569', dark='262126')
FONT = 'Microsoft YaHei'
entries = []

def col(value):
    return RGBColor.from_string(C.get(value, value))

def rect(s, x, y, w, h, fill='paper', line=None, radius=False):
    shape = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
                              Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid(); shape.fill.fore_color.rgb = col(fill)
    if line:
        shape.line.color.rgb = col(line); shape.line.width = Pt(1)
    else:
        shape.line.fill.background()
    if radius:
        shape.adjustments[0] = .12
    # Office's default shape style may otherwise add theme effects.
    shape._element.spPr.append(OxmlElement('a:effectLst'))
    return shape

def text(s, x, y, w, h, content, size=20, color='ink', bold=False, align=None):
    box = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear(); tf.word_wrap = True; tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.margin_left = tf.margin_right = 0
    tf.margin_top = tf.margin_bottom = 0
    for i, line in enumerate(content.split('\n')):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line; p.font.name = FONT; p.font.size = Pt(size)
        p.font.bold = bold; p.font.color.rgb = col(color)
        p.space_after = Pt(5); p.line_spacing = 1.15
        if align is not None: p.alignment = align
        for run in p.runs:
            rpr = run._r.get_or_add_rPr()
            ea = OxmlElement('a:ea'); ea.set('typeface', FONT); rpr.append(ea)
    return box

def line(s, x1, y1, x2, y2, color='red', width=2, arrow=False, dash=False):
    sh = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    sh.line.color.rgb = col(color); sh.line.width = Pt(width)
    if arrow:
        tail = OxmlElement('a:tailEnd'); tail.set('type', 'triangle'); sh.line._get_or_add_ln().append(tail)
    return sh

def tag(s, x, y, label, color='red', w=1.15):
    rect(s, x, y, w, .32, color)
    text(s, x+.07, y+.035, w-.14, .23, label, 11, 'white', True)

def node(s, x, y, w, label, fill='red', h=.73, size=20):
    rect(s, x, y, w, h, fill, radius=True)
    text(s, x+.1, y+.16, w-.2, h-.2, label, size, 'white' if fill != 'gold' else 'wine', True, PP_ALIGN.CENTER)

def footer(s, i, dark=False):
    color = 'light' if dark else 'muted'
    line(s, .55, 7.02, 12.78, 7.02, 'pink' if dark else 'line', .7)
    text(s, .57, 7.13, 11.3, .2, '人工智能  /  企业股权知识表示与风险推理  /  合成案例', 9, color)
    text(s, 12.13, 7.10, .64, .24, f'{i:02d} / 14', 10, color, False, PP_ALIGN.RIGHT)

def slide(section, title, sub, short, detail, sources, dark=False):
    s = P.slides.add_slide(P.slide_layouts[6])
    s.background.fill.solid(); s.background.fill.fore_color.rgb = col('wine' if dark else 'bg')
    rect(s, 0, 0, .15, 7.5, 'gold' if dark else 'red')
    text(s, .57, .32, 11, .3, section, 12, 'gold' if dark else 'red', True)
    text(s, .55, .9, 12.2, .67, title, 31, 'white' if dark else 'ink', True)
    text(s, .58, 1.7, 12, .65, sub, 17, 'light' if dark else 'muted')
    footer(s, len(P.slides), dark)
    note = '【3分钟精讲】\n'+short+'\n\n【展开讲解】\n'+detail+'\n\n【来源与口径】\n'+'\n'.join(sources)
    if len(P.slides) == 1: note += '\n\n组员与分工\n'+team_notes
    s.notes_slide.notes_text_frame.text = note
    entries.append(dict(slide=len(P.slides), title=title, narration=short, detail=detail, sources=sources))
    return s

def takeaway(s, label, dark=False):
    rect(s, .58, 6.32, 12.18, .48, 'gold' if dark else 'red')
    text(s, .79, 6.41, 11.77, .31, label, 16, 'wine' if dark else 'white', True)

def metric(s, x, y, value, label, w=2.85, dark=False):
    rect(s, x, y, w, 1.14, 'red' if dark else 'paper')
    text(s, x+.16, y+.1, w-.32, .58, value, 31, 'gold' if dark else 'red', True)
    text(s, x+.16, y+.75, w-.32, .28, label, 13, 'light' if dark else 'muted')

# 01 — Cover, with editable relationship motif.
s = slide('知识表示与推理  /  EQUITY KNOWLEDGE GRAPH', '企业股权知识表示与风险推理',
    '连接分散记录，追踪控制关系，为每条风险线索保留依据。',
    '企业背后还有企业，谁在控制、权益多少、哪些交易要核查？我们用股权案例展示三种知识表示怎样协同推理。',
    '本项目用16个虚构实体构建股权、担保、交易和亲属关系。汇报沿着背景、建模、案例推理和验证展开。控制关系、经济持股与风险标签具有不同语义，下面分别解释。全组共同参与开发、测试与汇报；完整姓名、学号和分工保留在本页备注。',
    ['项目报告（修订版）；dataset.py；组员信息沿用已发布修订版PPT。'], True)
text(s, .7, 2.66, 7.7, 1.6, '看清关系\n更要说清推理依据', 39, 'white', True)
text(s, .75, 4.43, 6.7, .53, '图谱找路径  ·  规则推结论  ·  框架存档案', 22, 'gold', True)
for a,b in [((9.0,3.0),(11.4,3.6)),((9.0,3.0),(9.5,5.1)),((11.4,3.6),(9.5,5.1)),((11.4,3.6),(12.1,5.7))]:
    line(s,*a,*b,'pink',3)
for x,y,label,fill in [(8.3,2.55,'自然人','gold'),(10.72,3.2,'企业 A','red'),(8.8,4.7,'企业 B','red'),(11.45,5.3,'企业 C','orange')]:
    node(s,x,y,1.32,label,fill,size=17)
text(s, .75, 5.68, 7.35, .6, '梁镇贤（组长） · 陈宗波 · 温文辉 · 陈宇冲 · 刘亦达', 17, 'white')
text(s, .75, 6.35, 7.2, .3, '项目团队  |  共同开发、测试与成果汇报', 13, 'light')

# 02 — Business context + carefully scoped real-world reference.
s=slide('01 / 背景与任务','贷款审批时，为什么一张股东表不够？',
    '目标：从“谁持股”进一步回答“谁控制、怎样关联、依据何在”。',
    '贷款审批和投资尽调需要穿透多层关系。ICIJ离岸数据库提供了关系追踪的真实行业参考；本项目用虚构数据演示建模、查询和证据核查。',
    '股东可能仍是企业，交易和担保又记录在不同表中。ICIJ Offshore Leaks公开数据库连接人员、企业和中介，启发我们以网络组织信息。本项目仅借鉴关系追踪思路，未导入其真实人员数据，也不声称复现其内部推理。关系出现本身不等于违法。',
    ['ICIJ, About the Offshore Leaks Database: https://offshoreleaks.icij.org/pages/about','项目报告（修订版）：背景与目标。'])
for x,n,title,body in [(.6,'01','多层股东','自然人 → 控股公司 → 子公司\n直接股东之外，还存在控制链。'),(4.73,'02','关联交易','交易双方名称不同\n背后却可能由同一人控制。'),(8.86,'03','交叉担保','多家企业彼此担保\n单条记录难以暴露闭合链路。')]:
    rect(s,x,2.62,3.85,1.83,'paper'); text(s,x+.18,2.78,.6,.48,n,24,'orange',True)
    text(s,x+.88,2.81,2.8,.45,title,23,'red',True); text(s,x+.2,3.48,3.46,.75,body,17)
rect(s,.6,4.7,12.11,1.32,'ink')
text(s,.8,4.88,3.25,.8,'行业参考\nICIJ Offshore Leaks',18,'gold',True)
text(s,4.25,4.91,8.05,.76,'连接人员、企业与中介，让关系可追踪。\n处理流程：记录整理 → 关系建模 → 查询推理 → 核查证据。',19,'white')
takeaway(s,'研究问题：让分散信息可连接，让推导过程可解释。')

# 03 — Data and processing pipeline.
s=slide('02 / 数据与系统流程','先定义实体与关系，再把记录变成知识',
    '同一批合成数据进入图谱、事实库与企业框架，避免三个互不相干的演示。',
    '数据包含四名自然人、十二家企业、二十五条关系。关系记录转换成五十三条初始事实，推理结果再写回图谱和企业档案。',
    '17条持股、4条担保、3条交易、1条亲属关系合计25条原始关系。53条初始事实还包括Person、Company与IsA等类型事实，所以事实数与图的边数不是同一口径。持股比例、交易金额、担保金额和亲属类型是关系属性。',
    ['dataset.py: demo_dataset；data/demo/*.csv','kr/knowledge_graph.py: to_predicates；main.py。'])
for x,v,l in [(.6,'4 + 12','自然人 + 企业'),(3.72,'25','原始关系'),(6.84,'53','初始谓词事实'),(9.96,'15','产生式规则')]:metric(s,x,2.51,v,l,w=2.76)
labels=[('HOLDS','17条 · 持股比例'),('GUARANTEES','4条 · 担保金额'),('TRADES','3条 · 交易金额'),('KINSHIP','1条 · 亲属类型')]
for i,(a,b) in enumerate(labels):
    x=.64+i*3.1; text(s,x,3.97,2.9,.35,a,16,'red',True);text(s,x,4.39,2.9,.39,b,17)
for i,(a,b) in enumerate([('CSV 记录','统一实体ID'),('知识表示','图谱 / 谓词 / 框架'),('推理与查询','规则链 / 路径 / 环路'),('知识增强','新增关系 / 档案更新')]):
    x=.64+i*3.1;rect(s,x,5.05,2.7,.94,'red' if i==2 else 'ink');text(s,x+.14,5.16,2.42,.35,a,20,'white',True);text(s,x+.14,5.61,2.42,.24,b,12,'gold')
    if i<3:line(s,x+2.74,5.49,x+3.01,5.49,'orange',2,True)
takeaway(s,'边、事实、推理结论分别计数；实体ID把不同知识表示连接起来。')

# 04 — Method comparison.
s=slide('03 / 方法比较','三种表示分别解决什么问题？',
    '选表示方法，要同时看表达能力、推理任务与维护代价。',
    '图谱擅长路径和环路；产生式用明确条件推导并解释结论；框架通过槽、继承和覆盖组织企业档案。三者互补，也各有维护成本。',
    '图谱将实体与带属性的关系连接起来，便于多跳查询，但依赖实体对齐。产生式使用谓词事实和IF–THEN规则，便于保存依据，但规则多时匹配与冲突处理成本增加。框架通过槽和继承管理属性，但缺省值不是已验证事实，全局路径计算仍需图查询。',
    ['kr/knowledge_graph.py；kr/production.py；kr/frames.py；reasoning/rules.py','参考资料：知识表示与知识图谱。'])
xs=[.6,2.53,6.08,9.35];ws=[1.93,3.55,3.27,3.37]
for x,w,head in zip(xs,ws,['表示','同一案例的表达','适合的任务','代价与局限']):
    rect(s,x,2.56,w,.57,'red');text(s,x+.15,2.7,w-.3,.31,head,17,'white',True)
rows=[('知识图谱','王建国 —持股80%→ 华鑫控股\n节点 + 有向边 + 属性','上游股东路径、循环关系\n把分散实体连接起来','实体同名需要核对\n关系时间与来源需治理'),('谓词 + 产生式','Holds(P1,C1,0.8)\nIF 持股 > 50% THEN 控制','控制传递、风险触发\n每条结论记录规则和前提','规则匹配与冲突消解\n简化规则需明确适用范围'),('框架表示','恒通科技 is_a 上市公司\n槽：控制人 / 风险 / 披露要求','档案组织、默认值继承\n推理后覆盖实例槽','默认低风险不代表安全\n不擅长全局路径计算')]
for i,row in enumerate(rows):
    y=3.13+i*.96
    for j,(x,w,val) in enumerate(zip(xs,ws,row)):
        rect(s,x,y,w,.95,'paper' if i%2==0 else 'light');text(s,x+.15,y+.16,w-.3,.7,val,17 if j==0 else 15,'red' if j==0 else 'ink',j==0)
takeaway(s,'图谱负责“找关系”，规则负责“推结论”，框架负责“组织与更新属性”。')

# 05 — Case diagram: share ratios as actual editable connectors.
s=slide('04 / 贯穿案例','恒通科技：同一条股权链，两个不同问题',
    '箭头方向：股东 → 被投资企业。名称和比例均来自本项目合成数据。',
    '王建国持有华鑫控股百分之八十；华鑫控股持有华鑫实业百分之六十五。两家公司分别持有恒通科技百分之二十和三十五。控制与经济权益要分开计算。',
    '图中只抽取恒通科技相关的4个节点与4条持股边，保留原始比例。Q1问能否按本项目规则形成控制关系；Q2问王建国沿两条路径对应的经济持股。前者由控制规则组合，后者为路径乘积再求和。节点简称对应CSV中的公司全称。',
    ['data/demo/holdings.csv；persons.csv；companies.csv。'])
line(s,2.86,3.11,4.07,3.11,'red',2.5,True);line(s,6.65,3.11,9.06,3.11,'red',2.5,True)
line(s,5.38,3.54,6.65,5.07,'orange',2.5,True);line(s,10.36,3.54,8.38,5.07,'orange',2.5,True)
node(s,.7,2.72,2.16,'王建国','ink');node(s,4.08,2.72,2.57,'华鑫控股');node(s,9.08,2.72,2.57,'华鑫实业')
node(s,6.13,5.1,2.8,'恒通科技','orange')
for x,y,w,t in [(3.0,2.55,.9,'80%'),(7.28,2.55,1.2,'65%'),(4.64,4.13,1.1,'20%'),(10.05,4.2,1.1,'35%')]:text(s,x,y,w,.42,t,24,'red',True)
tag(s,.74,4.52,'Q1 控制',w=1.12);text(s,.74,4.99,4.9,.38,'按当前规则，谁控制恒通科技？',20,'ink',True)
tag(s,.74,5.49,'Q2 权益',color='orange',w=1.12);text(s,2.04,5.47,3.8,.42,'王建国的经济持股是多少？',17)
takeaway(s,'先明确问题：控制关系的推导，与经济持股的计算，是两个任务。')

# 06 — Deep red proof page.
s=slide('05 / 控制关系推导','从 65% 到 55%：控制关系如何逐层成立？',
    '对应规则 R1 → R4 → R5 → R6；保留每一步的前提，才能解释最终结论。',
    '先按R1得到控股公司对子公司的控制，再按R4合并可支配的百分之二十和三十五，超过一半。R5传递控制，R6得到自然人控制人王建国。',
    'R1要求直接持股严格大于50%，所以恰好50%不成立。R4仅实现“自身与一个已控制子公司”合计持股过半的情形，并非任意多个子公司的通用合并。王建国对华鑫控股的80%满足R1；随后通过R5和R6得到当前模型下恒通科技的实际控制人结论。这里的55%是华鑫控股一侧可支配持股，不是王建国的经济权益。',
    ['reasoning/rules.py：R1、R4、R5、R6；output/demo_run.txt。'],True)
steps=[('R1','确立直接控制','华鑫控股 → 华鑫实业','65% > 50%'),('R4','合并可支配股份','华鑫控股 + 所控子公司','20% + 35% = 55%'),('R5','控制向上传递','王建国 → 华鑫控股 → 恒通','传递控制关系'),('R6','筛出自然人控制人','Controls + Person','王建国 → 恒通科技')]
for i,(r,t,b,v) in enumerate(steps):
    y=2.54+i*.85;tag(s,.72,y+.1,r,color='orange',w=.58);text(s,1.55,y+.09,3.05,.38,t,21,'white',True)
    text(s,4.74,y+.1,4.67,.43,b,19,'light');text(s,9.58,y+.07,2.85,.53,v,20,'gold',True)
    if i<3:line(s,1.02,y+.55,1.02,y+.88,'pink',1.7,True)
text(s,.8,6.0,11.7,.25,'R4覆盖本案例的“自身 + 单个受控子公司”；全部控制结论遵循当前模型的简化假设。',12,'light')
takeaway(s,'55% 回答“谁能支配”，不直接回答“王建国享有多少经济权益”。',True)

# 07 — Economic ownership.
s=slide('06 / 股权穿透计算','两条路径相乘再求和，经济持股为 34.2%',
    '路径内乘比例，路径间加贡献；把算式和路径同时展示，结果才可核查。',
    '经济持股按路径计算：百分之八十乘二十，得到十六；再乘六十五和三十五，得到十八点二。两条路径合计百分之三十四点二。',
    '第一条路径是王建国→华鑫控股→恒通科技，贡献16%；第二条经过华鑫实业，贡献18.2%。两条合计34.2%。本项目筛查以25%为阈值，且仅自然人可成为候选，并不构成法律认定。循环股权按简单路径截断，没有计算无限路径权益，因此结果边界必须说明。',
    ['reasoning/graph_reasoning.py: penetrate；output/demo_run.txt；test_project.py。'])
for y,n,path,eq in [(2.61,'路径 A','王建国 → 华鑫控股 → 恒通科技','80% × 20% = 16.0%'),(4.15,'路径 B','王建国 → 华鑫控股 → 华鑫实业 → 恒通科技','80% × 65% × 35% = 18.2%')]:
    rect(s,.66,y,8.23,1.28,'paper');tag(s,.87,y+.14,n,w=.88);text(s,2.0,y+.13,6.65,.41,path,17)
    text(s,.89,y+.7,7.62,.49,eq,27,'red',True)
rect(s,9.19,2.61,3.48,2.82,'red');text(s,9.47,2.92,2.93,.4,'经济持股合计',19,'white',True)
text(s,9.46,3.63,3.0,.81,'34.2%',44,'gold',True);text(s,9.48,4.72,2.98,.43,'16.0% + 18.2%',20,'white')
text(s,.84,5.8,11.7,.32,'筛查阈值设为25%；只有自然人可标记为候选，深度截断处的企业不能替代自然人。',14,'muted')
takeaway(s,'控制结论来自规则，经济权益来自路径；同一案例可以同时“受控”且“经济持股不足50%”。')

# 08 — Inference mechanism.
s=slide('07 / 推理机怎样工作','正向推理：新事实不断进入工作记忆',
    '匹配 → 冲突消解 → 执行并记录依据 → 检查是否还能产生新事实。',
    '推理机从已知事实出发，匹配规则，按优先级选择并执行，再记录前提。直到没有新事实，五十三条事实扩展为一百三十五条。',
    '综合数据库采用集合去重，并按谓词及首参数建立索引。每轮选最高优先级规则，同级按规则库顺序，执行所选规则在当前冲突集中的全部实例。不是每周期只产生一个事实。why保存规则与前提，trace保存过程。17周期后到达不动点，再运行无新增；现实现没有自动撤销过期推理结论。',
    ['kr/production.py: WorkingMemory、InferenceEngine.run、explain；reasoning/rules.py。'])
for i,(n,t,b) in enumerate([('01','匹配','寻找满足全部前提\n的变量绑定'),('02','冲突消解','优先级较高者先执行\n同级按规则库顺序'),('03','执行与解释','新事实加入工作记忆\n保存规则号和前提'),('04','终止判断','没有可产生新事实\n的规则时停止')]):
    x=.64+i*3.1;rect(s,x,2.68,2.73,1.72,'red' if i==1 else 'ink')
    text(s,x+.17,2.82,.52,.36,n,17,'gold',True);text(s,x+.79,2.8,1.9,.38,t,21,'white',True)
    text(s,x+.17,3.52,2.39,.66,b,16,'white')
    if i<3:line(s,x+2.76,3.5,x+3.04,3.5,'orange',2,True)
rect(s,.65,4.72,7.05,1.25,'paper');text(s,.88,4.92,6.59,.35,'控制识别 3  >  关联 / 可达 2  >  预警 1  >  定级 0',16,'red',True)
text(s,.88,5.43,6.4,.3,'优先级对应规则库；工作记忆去重，依据保存在 why 中。',14)
text(s,8.18,4.72,4.5,.68,'53 → 135',36,'red',True);text(s,8.24,5.53,4.24,.31,'17 个执行周期 · 新增 82 条事实',16,'muted')
takeaway(s,'“17周期”是规则执行轮数；“82条新增事实”是去重后的结论数。')

# 09 — Related transaction proof tree.
s=slide('08 / 风险证据链','一条交易预警，是怎样被推导出来的？',
    '以“远航贸易 → 恒通科技，5000万元”的合成交易为例，逐项核对前提。',
    '远航贸易和恒通科技按当前规则由王建国共同控制，再结合五千万元交易，推出关联交易线索。恒通属于上市公司，触发待核查预警；其中配偶一致行动是假设。',
    '远航控制链依赖配偶一致行动简化假设R2，再由R3合计王建国30%与李秀英25%，由持股较多者形成控制关系。恒通控制来自R1/R4/R5。R6取得两边自然人控制人，R7结合交易得到RelatedTrade，R12结合上市公司类型触发预警。实务还要核实一致行动证据、交易是否真实、是否已披露及交易定价，不能直接据此判违法。',
    ['data/demo/trades.csv；holdings.csv；kinships.csv','reasoning/rules.py: R2/R3/R6/R7/R12；output/demo_run.txt。'],True)
for x,t,b in [(.66,'事实 A：同一控制人','王建国 → 远航贸易\n王建国 → 恒通科技'),(4.8,'事实 B：存在交易','远航贸易 → 恒通科技\n金额：5000 万元'),(8.94,'事实 C：企业类型','恒通科技\nis_a 上市公司')]:
    rect(s,x,2.52,3.72,1.43,'red');text(s,x+.18,2.68,3.36,.33,t,19,'gold',True);text(s,x+.18,3.12,3.32,.81,b,19,'white')
line(s,2.52,3.91,4.65,4.46,'gold',2,True);line(s,6.64,3.91,5.55,4.46,'gold',2,True)
node(s,2.44,4.51,5.49,'R7：共同控制 + 交易 → 关联线索','orange',size=19)
line(s,8.02,4.88,8.67,4.88,'gold',2,True);line(s,10.75,3.91,10.75,4.46,'gold',2,True)
node(s,8.73,4.51,3.91,'R12：关联交易待核查','gold',size=18)
text(s,.82,5.69,11.57,.39,'关键假设：配偶视为一致行动人（R2），30% + 25% > 50%（R3）；实务需另行取证。',15,'light')
takeaway(s,'预警意味着“需要核查”；披露情况、协议和交易背景仍需补充证据。',True)

# 10 — Two clear cycles rather than a tiny full network screenshot.
s=slide('09 / 环路检测','循环持股与担保圈：两类关系，两种核查重点',
    '先按关系类型构建子图，再寻找回到自身的闭合路径。金额单位：万元。',
    '股权子图找到鼎丰投资、鼎丰建设和瑞祥建材的循环；担保子图找到鼎丰建设、天润商贸和博远咨询的环路。两圈有重叠，风险企业要去重计数。',
    '循环持股边分别为60%、40%、30%；担保边分别为3000、2000、2500万元。规则R8/R9与R10/R11先推导持股和担保可达，再由R13/R14检测自可达。图算法用simple_cycles作为另一种实现，测试核对节点集。两圈各3家，鼎丰建设重叠，环路企业共5家，加恒通交易预警共6家。',
    ['data/demo/holdings.csv；guarantees.csv；reasoning/rules.py: R8–R14','reasoning/graph_reasoning.py: find_cycles；test_project.py: test_cycle_algorithms_agree。'])
for x,title,fill,names,weights in [(.65,'循环持股 / 权益解释难度增加','red',['鼎丰投资','鼎丰建设','瑞祥建材'],['60%','40%','30%']),
                                  (6.81,'担保圈 / 关注潜在传染路径','green',['鼎丰建设','天润商贸','博远咨询'],['3000','2000','2500'])]:
    rect(s,x,2.48,5.87,3.49,'paper');text(s,x+.2,2.63,5.47,.39,title,20,fill,True)
    pts=[(x+2.15,3.2),(x+3.72,4.84),(x+.56,4.84)]
    # Diagonal edges stay outside the label boxes; bottom edge is leftward.
    line(s,x+3.33,3.9,x+4.26,4.81,fill,2.2,True)
    line(s,x+3.67,5.2,x+2.2,5.2,fill,2.2,True)
    line(s,x+1.07,4.79,x+2.47,3.9,fill,2.2,True)
    for (nx,ny),label in zip(pts,names):node(s,nx,ny,1.58,label,fill,size=15)
    text(s,x+4.14,4.03,1.13,.32,weights[0],18,fill,True)
    text(s,x+2.38,5.41,1.1,.31,weights[1],18,fill,True)
    text(s,x+.5,4.03,1.17,.32,weights[2],18,fill,True)
takeaway(s,'风险企业去重：循环持股 3 家 + 担保圈 3 家 − 重叠 1 家 + 交易预警 1 家 = 6 家。')

# 11 — Frame writeback.
s=slide('10 / 知识增强','结论写回企业框架，形成可更新的档案',
    '以恒通科技为例：类型继承提供公共属性，实例槽保存推理后的具体结果。',
    '恒通科技继承上市公司的披露属性；推理后，控制人由未识别变成王建国，风险提示变成关联交易待核查。默认低风险不代表已核实安全。',
    '框架继承链为企业→上市公司→恒通科技。披露要求和监管机构来自父框架，控制人与风险提示由实例覆盖。R15把任一预警定为“高”，属于简化定级规则。图谱同时新增实际控制边，使推理结果还能被后续查询使用。当前系统未自动撤回已失效结论，来源与时间管理是后续改进点。',
    ['kr/frames.py: build_company_frames；main.py：知识增强','reasoning/rules.py: R15；output/demo_run.txt。'])
tag(s,.68,2.54,'继承链',w=1.05);text(s,1.96,2.49,10.4,.48,'企业  →  上市公司  →  恒通科技',25,'red',True)
for x,w,h in [(.66,3.16,'属性槽'),(3.82,3.61,'推理前'),(7.43,5.23,'推理后')]:
    rect(s,x,3.17,w,.48,'red');text(s,x+.17,3.26,w-.34,.29,h,17,'white',True)
rows=[('实际控制人','未识别（默认）','王建国（R6）'),('风险等级','低（默认）','高（R15：简化定级）'),('风险提示','无（默认）','关联交易待核查（R12）'),('关联交易披露','必须公告披露（继承）','沿用上市公司类型属性')]
for i,(a,b,c) in enumerate(rows):
    y=3.65+i*.52
    for x,w,val in [(.66,3.16,a),(3.82,3.61,b),(7.43,5.23,c)]:
        rect(s,x,y,w,.51,'paper' if i%2==0 else 'light');text(s,x+.17,y+.1,w-.34,.35,val,17,'red' if x==7.43 else 'ink',x==7.43)
text(s,.84,5.95,11.7,.28,'实例值优先于继承值；缺省值用于组织档案，不等于现实世界已经确认的事实。',14,'muted')
takeaway(s,'知识增强 = 推理出新事实 + 写回关系图谱 + 覆盖实例属性。')

# 12 — Evidence / testing, including negative examples.
s=slide('11 / 结果与验证','除了“能推出来”，还要检查“不该推的没推”',
    '2026-10-05 本地功能复测：11 项测试通过；以下结果可在源码与日志中对应核查。',
    '本地十一项功能测试通过，覆盖控制链、经济持股、环路、边界和负例。例如恰好百分之五十不触发控制，深度截断的企业也不能标为自然人候选。',
    '正例检查恒通和远航的控制关系、34.2%权益和预警集合。负例检查星辰电子没有被错误识别出实际控制人。边界检查50%不触发R1、截断企业不作为自然人候选。框架测试检查继承、覆盖和非法槽值；重复变量匹配测试避免错误绑定；不动点测试核对二次运行无新增事实；SQL/图查询对照测试检查结果多重集合。11项通过验证实现逻辑，不是业务准确率。',
    ['test_project.py：11项单元测试；2026-10-05 使用 D:/python/python.exe 本地复测通过。','output/demo_run.txt；kr/production.py；reasoning/graph_reasoning.py。'])
for x,v,l in [(.65,'82','新增事实 / 135 − 53'),(3.75,'9','自然人控制边'),(6.85,'6','风险企业 / 去重后'),(9.95,'11 / 11','功能测试通过')]:metric(s,x,2.51,v,l,w=2.73)
items=[('正例','控制关系、34.2%权益、环路节点集','结论与合成案例一致'),('负例','星辰电子：没有错误控制人结论','没有证据时不强推'),('边界','50%不控制；截断企业非自然人候选','避免边界误判'),('一致性','框架覆盖 / 匹配 / 不动点 / SQL对照','不同机制相互验证')]
for i,(a,b,c) in enumerate(items):
    y=3.99+i*.5;tag(s,.78,y,a,w=.77,color='red' if i%2==0 else 'orange');text(s,1.83,y-.01,7.37,.34,b,17);text(s,9.43,y,3.06,.32,c,15,'muted')
takeaway(s,'测试验证当前实现的逻辑与边界；真实业务的准确率和召回率仍需单独评估。')

# 13 — Native PowerPoint editable chart + historical measurement context.
s=slide('12 / 查询效率实验','查询越深，三种实现的耗时差异越明显',
    '引用仓库既有基准：同一无环数据、同一路径语义，先核对结果一致，再计时。',
    '历史基准固定四百三十个实体、七百九十二条边，查询二十个目标并重复九轮。八层时DFS约零点五六四毫秒，快于本次两种SQL实现；结论只适用于该实验。',
    '图中每点是9轮中每轮20个目标平均查询耗时的中位数，单位毫秒/目标；包含结果物化，不含建库与绘图。四个深度均先逐目标比较路径结果的多重集合，避免把去重或剪枝差异当成性能优势。8层数据为DFS 0.563595、CTE 1.03337、JOIN 0.77047毫秒。来自原基准，不是本次重新测量，也不能外推为图数据库普遍更快。原环境Python 3.13.1、NetworkX 3.5、SQLite 3.45.3，seed=42。',
    ['output/benchmark.json、benchmark.csv；benchmark.py；kr/relational.py。'])
data=CategoryChartData();data.categories=['2层','4层','6层','8层']
methods=['图 DFS','SQL 递归 CTE','SQL 逐层 JOIN']
colors=['red','orange','green']
for method in methods:data.add_series(method,[next(v['median_ms'] for v in benchmark['results'] if v['method']==method and v['depth']==d) for d in [2,4,6,8]])
chart=s.shapes.add_chart(XL_CHART_TYPE.LINE_MARKERS, Inches(.8), Inches(2.66), Inches(7.65), Inches(3.4), data).chart
chart.has_legend=True;chart.legend.position=XL_LEGEND_POSITION.BOTTOM;chart.legend.include_in_layout=False
chart.legend.font.name=FONT;chart.legend.font.size=Pt(12)
chart.chart_style=13
chart.value_axis.minimum_scale=0;chart.value_axis.maximum_scale=1.2;chart.value_axis.major_unit=.3
chart.value_axis.tick_labels.number_format='0.0';chart.value_axis.tick_labels.font.size=Pt(11)
chart.category_axis.tick_labels.font.name=FONT;chart.category_axis.tick_labels.font.size=Pt(13)
chart.value_axis.has_major_gridlines=True
chart.value_axis.major_gridlines.format.line.color.rgb=col('line')
for series,c in zip(chart.series,colors):
    series.format.line.color.rgb=col(c);series.format.line.width=Pt(2.4)
    series.marker.size=7;series.marker.format.fill.solid();series.marker.format.fill.fore_color.rgb=col(c);series.marker.format.line.color.rgb=col(c)
text(s,.87,2.33,7.4,.25,'耗时 / 毫秒 / 目标（9轮中位数）',12,'muted')
rect(s,8.8,2.61,3.84,3.38,'ink');text(s,9.05,2.8,3.35,.38,'8 层查询',22,'white',True)
for i,(m,c) in enumerate(zip(methods,colors)):
    val=next(v['median_ms'] for v in benchmark['results'] if v['method']==m and v['depth']==8)
    text(s,9.08,3.38+i*.68,2.8,.28,m,14,'light');text(s,9.08,3.65+i*.68,3.0,.4,f'{val:.3f} ms',24,'gold',True)
text(s,.84,6.05,11.8,.22,'430实体 · 792条边 · 20个目标 · 9轮重复；包含结果物化，不含建库与绘图。',12,'muted')
takeaway(s,'这里只能比较当前数据与实现；“图查询普遍优于SQL”不是本实验能够支持的结论。')

# 14 — Closing and grounded roadmap.
s=slide('13 / 总结与改进','下一步：增强数据可信度与推理可靠性',
    '已实现知识表示、推理解释与结果回写；下一步完善数据、规则和更新机制。',
    '我们已实现知识表示、推理解释与结果回写。下一步补足身份与来源核验、控制协议、时间变化和失效结论撤回，让每条线索不仅算得出，也查得到依据。',
    '数据方面，加入身份对齐、来源凭证和时间戳；规则方面，从配偶等简化假设扩展到经验证的控制协议、投票权约束和规则冲突处理；算法方面，改进循环持股权益求解、撤回依赖失效事实的结论，并扩大基准规模。本项目展示三种知识表示的任务适配，不提供法律判断或生产环境准确率保证。完整团队分工见第1页备注，完整参考文献见修订报告。',
    ['项目报告（修订版）：问题与改进、总结、参考文献。','ICIJ Offshore Leaks: https://offshoreleaks.icij.org/pages/about','Hogan et al., Knowledge Graphs, ACM Computing Surveys, 2021, doi:10.1145/3447772','NetworkX: https://networkx.org/documentation/stable/；SQLite: https://www.sqlite.org/lang_with.html'],True)
for x,tagline,title,body in [(.66,'DATA','可信数据','实体身份核验与来源凭证\n为关系补充时间与版本'),(4.8,'RULES','可信规则','核实控制协议与投票权\n明确适用条件与冲突处理'),(8.94,'UPDATES','可信更新','结论随证据失效而撤回\n改进循环权益与规模测试')]:
    rect(s,x,2.6,3.71,2.05,'red');text(s,x+.2,2.78,3.3,.28,tagline,13,'gold',True)
    text(s,x+.2,3.24,3.31,.44,title,25,'white',True);text(s,x+.2,3.91,3.3,.73,body,17,'light')
text(s,.8,5.04,11.8,.68,'让关系可计算，让结论有依据。',33,'white',True)
text(s,.81,5.94,11.7,.26,'参考：ICIJ Offshore Leaks · Hogan等《Knowledge Graphs》· NetworkX / SQLite文档',11,'light')
takeaway(s,'三种表示协同：图谱找路径 → 规则推结论 → 框架存档案 → 测试核对边界。',True)

assert len(P.slides)==14
P.core_properties.title='企业股权知识表示与风险推理｜暖色增强版'
P.core_properties.subject='企业股权分析项目；14页；酒红与琥珀金；可编辑图表与关系图'
P.core_properties.author='项目团队'
P.core_properties.keywords='知识图谱,产生式,框架,股权穿透,风险推理'
P.save(TARGET)
(QA/'content.json').write_text(json.dumps(entries,ensure_ascii=False,indent=2),encoding='utf-8')
talk='# 企业股权知识表示与风险推理：汇报讲稿（暖色增强版）\n\n对应14页PPT。每页分为精讲正文和可展开讲解；精讲正文供三分钟视频排练，实际时长需按语速实测。\n\n'
for e in entries:
    talk+=f"## 第{e['slide']}页 {e['title']}\n\n**精讲**\n\n{e['narration']}\n\n**展开讲解**\n\n{e['detail']}\n\n"
talk+='## 组员与分工\n\n'+team_notes+'\n'
(OUT/'汇报讲稿_暖色增强版.md').write_text(talk,encoding='utf-8')
print(json.dumps({'file':str(TARGET),'slides':len(P.slides),'narration_characters':sum(len(e['narration']) for e in entries)},ensure_ascii=False))
